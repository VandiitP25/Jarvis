from __future__ import annotations

import os
import re
import secrets
import subprocess
import time
import webbrowser

from pathlib import Path, PureWindowsPath
from threading import Lock
from urllib.parse import quote, urlparse
from typing import Literal

from fastapi import APIRouter, HTTPException, Query
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document


# ---------------------------------------------------------------------------
# Windows notifications
# ---------------------------------------------------------------------------

def send_jarvis_notification(title: str, message: str) -> None:
    """Show a Windows toast notification without interrupting automation."""
    try:
        from winotify import Notification

        toast = Notification(
            app_id="JARVIS",
            title=title,
            msg=message,
            duration="short",
        )
        toast.show()
    except Exception as exc:
        # A notification failure must not undo a completed action.
        print(f"JARVIS notification failed: {exc}")


router = APIRouter(prefix="/api/tools", tags=["JARVIS Tools"])


# ---------------------------------------------------------------------------
# Project and workspace configuration
# ---------------------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parent.parent
WORKSPACE = PROJECT_ROOT / "JARVIS_Workspace"
NOTES_DIR = WORKSPACE / "Notes"

WORKSPACE.mkdir(parents=True, exist_ok=True)
NOTES_DIR.mkdir(parents=True, exist_ok=True)

HOME = Path.home()
ONEDRIVE = Path(
    os.environ.get("OneDrive", str(HOME / "OneDrive"))
).expanduser()

APPROVED_CANDIDATES: dict[str, list[Path]] = {
    "Desktop": [ONEDRIVE / "Desktop", HOME / "Desktop"],
    "Documents": [ONEDRIVE / "Documents", HOME / "Documents"],
    "Downloads": [HOME / "Downloads", ONEDRIVE / "Downloads"],
    "Project": [PROJECT_ROOT],
}

READABLE_EXTENSIONS = {
    ".txt", ".md", ".py", ".js", ".html", ".css", ".json",
    ".csv", ".xml", ".yaml", ".yml", ".log", ".ini", ".toml",
    ".v", ".sv", ".c", ".cpp", ".h", ".java",
}

IGNORED_DIRECTORY_NAMES = {
    "node_modules", "__pycache__", ".git", ".venv", "venv",
}

MAX_READ_BYTES = 1_000_000
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_PDF_PAGES = 300
MAX_SEARCH_RESULTS = 500

# Only these applications can be launched by the automation API.
ALLOWED_APPS = {
    "notepad": ["notepad.exe"],
    "calculator": ["calc.exe"],
    "paint": ["mspaint.exe"],
}

# Temporary, single-use approval tokens.
# These disappear whenever the backend restarts.
PENDING_ACTIONS: dict[str, dict] = {}
PENDING_LOCK = Lock()
APPROVAL_EXPIRY_SECONDS = 120


# ---------------------------------------------------------------------------
# Request models
# ---------------------------------------------------------------------------

class LocalPathRequest(BaseModel):
    path: str = Field(min_length=1, max_length=1000)


class CreateNoteRequest(BaseModel):
    title: str = Field(min_length=1, max_length=100)
    content: str = Field(min_length=1, max_length=20000)


class AutomationRequest(BaseModel):
    action: Literal["open_app", "open_website", "open_file"]
    target: str = Field(min_length=1, max_length=1000)


class AutomationApprovalRequest(BaseModel):
    approval_id: str = Field(min_length=1, max_length=100)
    approved: bool


# ---------------------------------------------------------------------------
# Approved roots and secure path resolution
# ---------------------------------------------------------------------------

def available_roots() -> dict[str, Path]:
    """Return existing approved directories."""
    roots: dict[str, Path] = {}

    for name, candidates in APPROVED_CANDIDATES.items():
        for candidate in candidates:
            try:
                resolved = candidate.expanduser().resolve()
                if resolved.is_dir():
                    roots[name] = resolved
                    break
            except (OSError, RuntimeError):
                continue

    roots["JARVIS_Workspace"] = WORKSPACE.resolve()
    return roots


def resolve_local_path(path_text: str) -> Path:
    """
    Resolve a relative path such as:
    Desktop/Research Papers/paper.pdf

    Absolute paths and paths escaping approved roots are rejected.
    """
    if not isinstance(path_text, str):
        raise HTTPException(
            status_code=400,
            detail="Path must be a string.",
        )

    cleaned = path_text.strip().replace("\\", "/")

    if not cleaned:
        raise HTTPException(
            status_code=400,
            detail="Enter a file path.",
        )

    if cleaned.lower().startswith(("file:", "http:", "https:")):
        raise HTTPException(
            status_code=400,
            detail="Use a relative path, not a URL.",
        )

    windows_path = PureWindowsPath(path_text.strip())

    if (
        windows_path.is_absolute()
        or windows_path.drive
        or cleaned.startswith("/")
    ):
        raise HTTPException(
            status_code=400,
            detail="Use a relative path returned by JARVIS, not an absolute path.",
        )

    parts = [part for part in cleaned.split("/") if part]

    if not parts or any(part in {".", ".."} for part in parts):
        raise HTTPException(
            status_code=403,
            detail="Invalid file path.",
        )

    roots = available_roots()
    root = roots.get(parts[0])

    if root is None:
        raise HTTPException(
            status_code=403,
            detail=(
                f"'{parts[0]}' is not an available approved folder. "
                "Check GET /api/tools/locations."
            ),
        )

    try:
        target = root.joinpath(*parts[1:]).resolve(strict=False)
        target.relative_to(root)
    except (ValueError, OSError, RuntimeError):
        raise HTTPException(
            status_code=403,
            detail="Access outside approved folders is not allowed.",
        )

    return target


# ---------------------------------------------------------------------------
# List approved folders
# ---------------------------------------------------------------------------

@router.get("/locations")
def list_locations():
    return {
        "locations": [
            {"name": name, "path": str(path)}
            for name, path in available_roots().items()
        ]
    }


# ---------------------------------------------------------------------------
# Search local files by filename
# ---------------------------------------------------------------------------

@router.get("/files")
def list_files(
    folder: str = Query("Documents", min_length=1, max_length=50),
    q: str = Query("", max_length=200),
):
    roots = available_roots()
    root = roots.get(folder)

    if root is None:
        raise HTTPException(
            status_code=400,
            detail="Unknown or unavailable folder. Check GET /api/tools/locations.",
        )

    query = q.strip().casefold()
    results: list[dict] = []
    truncated = False

    for current_dir, dir_names, file_names in os.walk(
        root, topdown=True, followlinks=False
    ):
        current = Path(current_dir)

        dir_names[:] = [
            name for name in dir_names
            if not name.startswith(".")
            and name.lower() not in IGNORED_DIRECTORY_NAMES
            and not (current / name).is_symlink()
        ]

        for filename in file_names:
            item = current / filename

            try:
                if filename.startswith(".") or item.is_symlink():
                    continue

                if not item.is_file():
                    continue

                if query and query not in filename.casefold():
                    continue

                resolved = item.resolve()
                resolved.relative_to(root)
                relative = item.relative_to(root)

                results.append({
                    "name": item.name,
                    "path": f"{folder}/{relative.as_posix()}",
                    "type": item.suffix.lower() or "file",
                    "size_bytes": item.stat().st_size,
                })

                if len(results) >= MAX_SEARCH_RESULTS:
                    truncated = True
                    break

            except (OSError, ValueError, RuntimeError):
                continue

        if truncated:
            break

    return {
        "folder": folder,
        "query": q,
        "count": len(results),
        "files": results,
        "truncated": truncated,
    }


# ---------------------------------------------------------------------------
# Open a local file
# ---------------------------------------------------------------------------

@router.post("/open-file")
def open_file(request: LocalPathRequest):
    target = resolve_local_path(request.path)

    if not target.is_file():
        raise HTTPException(
            status_code=404,
            detail="File not found.",
        )

    if os.name != "nt":
        raise HTTPException(
            status_code=501,
            detail="This feature requires Windows.",
        )

    try:
        os.startfile(str(target))  # type: ignore[attr-defined]
    except OSError as exc:
        raise HTTPException(
            status_code=500,
            detail="Windows could not open this file.",
        ) from exc

    return {"message": "File opened.", "path": str(target)}


# ---------------------------------------------------------------------------
# Read local text files
# ---------------------------------------------------------------------------

@router.post("/read-file")
def read_file(request: LocalPathRequest):
    target = resolve_local_path(request.path)

    if not target.is_file():
        raise HTTPException(
            status_code=404,
            detail="File not found.",
        )

    if target.suffix.lower() not in READABLE_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail="Unsupported text file format.",
        )

    try:
        if target.stat().st_size > MAX_READ_BYTES:
            raise HTTPException(
                status_code=413,
                detail="File exceeds the 1 MB reading limit.",
            )

        content = target.read_text(encoding="utf-8-sig")

    except HTTPException:
        raise
    except UnicodeDecodeError as exc:
        raise HTTPException(
            status_code=415,
            detail="The file is not valid UTF-8 text.",
        ) from exc
    except OSError as exc:
        raise HTTPException(
            status_code=500,
            detail="Windows could not read this file.",
        ) from exc

    return {
        "name": target.name,
        "path": request.path,
        "content": content,
    }


# ---------------------------------------------------------------------------
# Read PDF and Word documents
# ---------------------------------------------------------------------------

@router.post("/read-document")
def read_document(request: LocalPathRequest):
    target = resolve_local_path(request.path)

    if not target.is_file():
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    extension = target.suffix.lower()

    if extension not in {".pdf", ".docx"}:
        raise HTTPException(
            status_code=415,
            detail="Supported document formats are PDF and DOCX.",
        )

    try:
        if target.stat().st_size > MAX_DOCUMENT_BYTES:
            raise HTTPException(
                status_code=413,
                detail="Document exceeds the 25 MB limit.",
            )

        if extension == ".pdf":
            reader = PdfReader(str(target), strict=False)

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=415,
                    detail="PDF is encrypted or password-protected.",
                )

            if len(reader.pages) > MAX_PDF_PAGES:
                raise HTTPException(
                    status_code=413,
                    detail="PDF exceeds the 300-page limit.",
                )

            pages: list[str] = []

            for page_number, page in enumerate(reader.pages, start=1):
                try:
                    page_text = page.extract_text() or ""
                except Exception as exc:
                    raise HTTPException(
                        status_code=422,
                        detail=(
                            f"Could not extract text from PDF page "
                            f"{page_number}. The PDF may be damaged."
                        ),
                    ) from exc

                pages.append(f"[Page {page_number}]\n{page_text}")

            content = "\n\n".join(pages)

        else:
            document = Document(str(target))
            sections: list[str] = []

            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    sections.append(paragraph.text)

            for table_index, table in enumerate(document.tables, start=1):
                sections.append(f"[Table {table_index}]")

                for row in table.rows:
                    sections.append(
                        " | ".join(cell.text.strip() for cell in row.cells)
                    )

            content = "\n".join(sections)

    except HTTPException:
        raise
    except (OSError, ValueError) as exc:
        raise HTTPException(
            status_code=422,
            detail="Could not open this document. It may be damaged or unsupported.",
        ) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="Could not extract text from this document.",
        ) from exc

    return {
        "name": target.name,
        "path": request.path,
        "characters": len(content),
        "text": content,
        "message": (
            "Text extracted locally."
            if content.strip()
            else (
                "No text extracted. The PDF may be scanned or contain "
                "images; OCR may be required."
            )
        ),
    }


# ---------------------------------------------------------------------------
# JARVIS notes
# ---------------------------------------------------------------------------

@router.get("/notes")
def list_notes():
    return {
        "notes": sorted(
            path.name
            for path in NOTES_DIR.glob("*.txt")
            if path.is_file()
        )
    }


@router.post("/notes")
def create_note(request: CreateNoteRequest):
    title = re.sub(
        r'[<>:"/\\|?*\x00-\x1f]',
        "",
        request.title,
    ).strip().rstrip(".")

    if not title or title in {".", ".."}:
        raise HTTPException(
            status_code=400,
            detail="Invalid note title.",
        )

    filename = title[:100] + ".txt"
    notes_root = NOTES_DIR.resolve()
    target = (notes_root / filename).resolve()

    try:
        target.relative_to(notes_root)
    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail="Invalid note path.",
        ) from exc

    if target.exists():
        raise HTTPException(
            status_code=409,
            detail="A note with that title already exists.",
        )

    try:
        target.write_text(request.content, encoding="utf-8")
    except OSError as exc:
        raise HTTPException(
            status_code=500,
            detail="Could not save the note.",
        ) from exc

    return {"message": "Note saved.", "title": filename}


@router.get("/notes/{title}")
def read_note(title: str):
    if Path(title).name != title or not title.endswith(".txt"):
        raise HTTPException(
            status_code=400,
            detail="Invalid note filename.",
        )

    notes_root = NOTES_DIR.resolve()
    target = (notes_root / title).resolve()

    try:
        target.relative_to(notes_root)
    except ValueError as exc:
        raise HTTPException(
            status_code=404,
            detail="Note not found.",
        ) from exc

    if not target.is_file():
        raise HTTPException(
            status_code=404,
            detail="Note not found.",
        )

    try:
        content = target.read_text(encoding="utf-8")
    except OSError as exc:
        raise HTTPException(
            status_code=500,
            detail="Could not read the note.",
        ) from exc

    return {"title": title, "content": content}


# ---------------------------------------------------------------------------
# Browser search
# ---------------------------------------------------------------------------

@router.get("/search")
def browser_search(q: str = Query(..., min_length=1, max_length=500)):
    query = q.strip()

    if not query:
        raise HTTPException(
            status_code=400,
            detail="Enter a search query.",
        )

    url = "https://www.google.com/search?q=" + quote(query)

    try:
        opened = webbrowser.open(url, new=2)
    except Exception:
        opened = False

    return {
        "query": query,
        "url": url,
        "browser_open_requested": bool(opened),
        "message": (
            "Search page requested in the default browser."
            if opened
            else "Open the returned URL in your browser."
        ),
    }


# ---------------------------------------------------------------------------
# Automation validation
# ---------------------------------------------------------------------------

def clean_expired_approvals() -> None:
    """Remove expired approvals. Caller must hold PENDING_LOCK."""
    now = time.time()

    expired = [
        token
        for token, item in PENDING_ACTIONS.items()
        if item["expires_at"] <= now
    ]

    for token in expired:
        PENDING_ACTIONS.pop(token, None)


def validate_automation_target(action: str, target: str) -> str:
    """Validate the target without performing the requested action."""
    target = target.strip()

    if not target:
        raise HTTPException(
            status_code=400,
            detail="Target cannot be empty.",
        )

    if action == "open_app":
        app_name = target.casefold()

        if app_name not in ALLOWED_APPS:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Application not allowed. Choose Notepad, "
                    "Calculator, or Paint."
                ),
            )

        return app_name

    if action == "open_website":
        parsed = urlparse(target)

        if (
            parsed.scheme not in {"http", "https"}
            or not parsed.hostname
            or parsed.username is not None
            or parsed.password is not None
            or any(ord(char) < 32 for char in target)
        ):
            raise HTTPException(
                status_code=400,
                detail="Enter a valid HTTP or HTTPS website URL.",
            )

        return target

    if action == "open_file":
        resolved = resolve_local_path(target)

        if not resolved.is_file():
            raise HTTPException(
                status_code=404,
                detail="File not found.",
            )

        return target

    raise HTTPException(
        status_code=400,
        detail="Unsupported automation action.",
    )


# ---------------------------------------------------------------------------
# Prepare an automation action (does not execute it)
# ---------------------------------------------------------------------------

@router.post("/automation/prepare")
def prepare_automation(request: AutomationRequest):
    target = validate_automation_target(
        request.action,
        request.target,
    )

    descriptions = {
        "open_app": f"Launch the application: {target}",
        "open_website": f"Open this website: {target}",
        "open_file": f"Open this approved file: {target}",
    }

    approval_id = secrets.token_urlsafe(24)
    expires_at = time.time() + APPROVAL_EXPIRY_SECONDS

    with PENDING_LOCK:
        clean_expired_approvals()

        PENDING_ACTIONS[approval_id] = {
            "action": request.action,
            "target": target,
            "expires_at": expires_at,
        }

    return {
        "status": "awaiting_approval",
        "approval_id": approval_id,
        "action": request.action,
        "target": target,
        "description": descriptions[request.action],
        "expires_in_seconds": APPROVAL_EXPIRY_SECONDS,
        "message": "Review the action and explicitly approve or reject it.",
    }


# ---------------------------------------------------------------------------
# Approve or reject a prepared action
# ---------------------------------------------------------------------------

@router.post("/automation/approve")
def approve_automation(request: AutomationApprovalRequest):
    # Remove the token immediately so it cannot be reused.
    with PENDING_LOCK:
        clean_expired_approvals()
        pending = PENDING_ACTIONS.pop(request.approval_id, None)

    if pending is None:
        raise HTTPException(
            status_code=404,
            detail=(
                "Approval not found, already used, or expired. "
                "Prepare the action again."
            ),
        )

    if not request.approved:
        return {
            "status": "rejected",
            "message": "Action rejected. Nothing was executed.",
        }

    action = pending["action"]

    # Validate again immediately before execution.
    target = validate_automation_target(
        action,
        pending["target"],
    )

    try:
        if action == "open_app":
            if os.name != "nt":
                raise HTTPException(
                    status_code=501,
                    detail="Application launching requires Windows.",
                )

            # The executable comes only from the fixed allowlist.
            subprocess.Popen(
                ALLOWED_APPS[target],
                shell=False,
            )

        elif action == "open_website":
            opened = webbrowser.open(target, new=2)

            if not opened:
                raise HTTPException(
                    status_code=500,
                    detail="The default browser could not open the website.",
                )

        elif action == "open_file":
            if os.name != "nt":
                raise HTTPException(
                    status_code=501,
                    detail="Opening files this way requires Windows.",
                )

            file_path = resolve_local_path(target)

            if not file_path.is_file():
                raise HTTPException(
                    status_code=404,
                    detail="File no longer exists.",
                )

            os.startfile(str(file_path))  # type: ignore[attr-defined]

        else:
            raise HTTPException(
                status_code=400,
                detail="Unsupported automation action.",
            )

    except HTTPException:
        raise
    except OSError as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Windows could not complete the action: {exc}",
        ) from exc

    action_description = {
        "open_app": f"Opened application: {target}",
        "open_website": f"Opened website: {target}",
        "open_file": f"Opened file: {target}",
    }.get(action, f"Completed action: {action}")

    send_jarvis_notification(
        "JARVIS - Task Completed",
        action_description,
    )

    return {
        "status": "completed",
        "action": action,
        "message": "The approved action was executed.",
    }